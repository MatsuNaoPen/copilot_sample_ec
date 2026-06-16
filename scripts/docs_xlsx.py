#!/usr/bin/env python3
"""要件定義書・設計書を Markdown ⇄ xlsx で双方向変換するツール。

設計方針（双方向＝ラウンドトリップ対応）:
  Markdown と xlsx を直接変換し合わず、間に「ブロックモデル」（見出し / 段落 /
  箇条書き / 番号付き / チェックリスト / 表 / コードブロック / 引用）を置く。

      docs/*.md  ⇄  [ブロックモデル]  ⇄  docs/*.xlsx
                  parse / render        write / read

  - Markdown を正本（source of truth）とし、git ではこちらを差分レビューする。
  - xlsx には A 列（非表示）に各行の「種別キー」を埋め込み、取り込み（import）は
    座標ではなくこのキーを基準に行う。Excel で行を増減してもキーが残る限り復元できる。
    ※Excel で行を増やすときは「既存行をコピー」して貼り付けると、非表示の A 列キーも
      一緒に複製されるため import で正しく取り込める。
  - 往復対象はブロックモデルで表現できる構造のみ（ロッシー）。Excel 上で足した色・
    セル結合・別シートなどは取り込まれない。

使い方:
  python3 scripts/docs_xlsx.py export <input.md> <output.xlsx>
  python3 scripts/docs_xlsx.py import <input.xlsx> <output.md>
  python3 scripts/docs_xlsx.py export-docs   # docs/{requirements,design}.md -> .xlsx
  python3 scripts/docs_xlsx.py import-docs    # docs/{requirements,design}.xlsx -> .md
  python3 scripts/docs_xlsx.py template       # docs/templates/*_template.md -> .xlsx
"""

from __future__ import annotations

import argparse
import os
import re
import sys
from dataclasses import dataclass, field
from typing import Optional

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:  # pragma: no cover - 環境セットアップ案内
    sys.stderr.write(
        "openpyxl が見つかりません。先に依存をインストールしてください:\n"
        "  python3 -m pip install -r scripts/requirements.txt\n"
    )
    sys.exit(2)


# ---------------------------------------------------------------------------
# ブロックモデル
# ---------------------------------------------------------------------------

@dataclass
class Block:
    """Markdown / xlsx 双方の中間表現となる 1 ブロック。

    kind ごとに使うフィールドが異なる:
      heading            : level(1-4), text
      paragraph / quote  : text（複数行は \n を保持）
      bullet / ordered   : level(0-), text
      task               : checked(bool), text
      code               : lang, text（複数行は \n を保持）
      table              : header(list[str]), rows(list[list[str]])
    """

    kind: str
    text: str = ""
    level: int = 1
    checked: bool = False
    lang: str = ""
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)


LIST_KINDS = {"bullet", "ordered", "task"}


# ---------------------------------------------------------------------------
# Markdown -> ブロックモデル
# ---------------------------------------------------------------------------

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
_TASK_RE = re.compile(r"^(\s*)[-*]\s+\[([ xX])\]\s+(.*)$")
_ORDERED_RE = re.compile(r"^(\s*)\d+\.\s+(.*)$")
_FENCE_RE = re.compile(r"^```(.*)$")


def _is_table_separator(line: str) -> bool:
    s = line.strip()
    if "|" not in s or set(s) - set("|:- \t") != set():
        return False
    return bool(re.fullmatch(r"\|?\s*:?-{1,}:?\s*(\|\s*:?-{1,}:?\s*)*\|?", s))


def _split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [cell.strip() for cell in s.split("|")]


def _indent_level(spaces: str) -> int:
    return len(spaces.replace("\t", "  ")) // 2


def parse_markdown(text: str) -> list[Block]:
    """Markdown テキストをブロックの並びへ。テンプレートで使う構文の部分集合に対応。"""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    while i < n:
        line = lines[i]

        # 空行はブロック区切り。
        if line.strip() == "":
            i += 1
            continue

        # コードフェンス ```lang ... ```
        m = _FENCE_RE.match(line)
        if m:
            lang = m.group(1).strip()
            body: list[str] = []
            i += 1
            while i < n and not _FENCE_RE.match(lines[i]):
                body.append(lines[i])
                i += 1
            i += 1  # 終端フェンスを読み飛ばす
            blocks.append(Block("code", text="\n".join(body), lang=lang))
            continue

        # 見出し
        m = _HEADING_RE.match(line)
        if m:
            level = len(m.group(1))
            blocks.append(Block("heading", text=m.group(2).strip(), level=level))
            i += 1
            continue

        # 表（ヘッダ行 + 区切り行 + データ行）
        if "|" in line and i + 1 < n and _is_table_separator(lines[i + 1]):
            header = _split_table_row(line)
            i += 2
            rows: list[list[str]] = []
            while i < n and "|" in lines[i] and lines[i].strip() != "":
                if _is_table_separator(lines[i]):
                    i += 1
                    continue
                cells = _split_table_row(lines[i])
                # 列数をヘッダに合わせて整える。
                if len(cells) < len(header):
                    cells += [""] * (len(header) - len(cells))
                rows.append(cells[: len(header)])
                i += 1
            blocks.append(Block("table", header=header, rows=rows))
            continue

        # チェックリスト（- [ ] / - [x]）
        m = _TASK_RE.match(line)
        if m:
            checked = m.group(2).lower() == "x"
            blocks.append(
                Block("task", text=m.group(3).strip(),
                      level=_indent_level(m.group(1)), checked=checked)
            )
            i += 1
            continue

        # 箇条書き（- / *）
        m = _BULLET_RE.match(line)
        if m:
            blocks.append(
                Block("bullet", text=m.group(2).strip(), level=_indent_level(m.group(1)))
            )
            i += 1
            continue

        # 番号付き（1. 2. ...）
        m = _ORDERED_RE.match(line)
        if m:
            blocks.append(
                Block("ordered", text=m.group(2).strip(), level=_indent_level(m.group(1)))
            )
            i += 1
            continue

        # 引用（> ...）連続行をまとめる
        if line.lstrip().startswith(">"):
            quote: list[str] = []
            while i < n and lines[i].lstrip().startswith(">"):
                stripped = lines[i].lstrip()[1:]
                if stripped.startswith(" "):
                    stripped = stripped[1:]
                quote.append(stripped)
                i += 1
            blocks.append(Block("quote", text="\n".join(quote)))
            continue

        # それ以外は段落。次の空行/構造行まで連結（HTML コメント等もそのまま保持）。
        para: list[str] = []
        while i < n and lines[i].strip() != "":
            cur = lines[i]
            if (
                _HEADING_RE.match(cur)
                or _FENCE_RE.match(cur)
                or _BULLET_RE.match(cur)
                or _ORDERED_RE.match(cur)
                or cur.lstrip().startswith(">")
                or ("|" in cur and i + 1 < n and _is_table_separator(lines[i + 1]))
            ):
                break
            para.append(cur)
            i += 1
        blocks.append(Block("paragraph", text="\n".join(para)))

    return blocks


# ---------------------------------------------------------------------------
# ブロックモデル -> Markdown
# ---------------------------------------------------------------------------

def render_markdown(blocks: list[Block]) -> str:
    """ブロックの並びを正規化された Markdown へ。往復で安定するよう整形する。"""
    out: list[str] = []
    prev: Optional[Block] = None
    ordered_counter = 0
    ordered_level = -1

    for b in blocks:
        # ブロック間の空行ルール: 連続するリスト項目どうしは詰める。
        if prev is not None:
            both_list = prev.kind in LIST_KINDS and b.kind in LIST_KINDS
            if not both_list:
                out.append("")

        if b.kind == "heading":
            out.append("#" * b.level + " " + b.text)
        elif b.kind == "paragraph":
            out.extend(b.text.split("\n"))
        elif b.kind == "quote":
            out.extend("> " + ln if ln else ">" for ln in b.text.split("\n"))
        elif b.kind == "code":
            out.append("```" + b.lang)
            if b.text:
                out.extend(b.text.split("\n"))
            out.append("```")
        elif b.kind == "bullet":
            out.append("  " * b.level + "- " + b.text)
        elif b.kind == "task":
            box = "[x]" if b.checked else "[ ]"
            out.append("  " * b.level + f"- {box} " + b.text)
        elif b.kind == "ordered":
            if prev is not None and prev.kind == "ordered" and prev.level == b.level:
                ordered_counter += 1
            else:
                ordered_counter = 1
                ordered_level = b.level
            out.append("  " * b.level + f"{ordered_counter}. " + b.text)
        elif b.kind == "table":
            out.append("| " + " | ".join(b.header) + " |")
            out.append("|" + "|".join("------" for _ in b.header) + "|")
            for row in b.rows:
                cells = list(row) + [""] * (len(b.header) - len(row))
                out.append("| " + " | ".join(cells[: len(b.header)]) + " |")

        prev = b

    text = "\n".join(out).strip("\n")
    return text + "\n"


# ---------------------------------------------------------------------------
# スタイル定義（アプリのテーマ #4c6ef5 に合わせる）
# ---------------------------------------------------------------------------

KEY_COL = 1          # A 列: 非表示の種別キー
CONTENT_COL = 2      # B 列: 内容の開始列
DEFAULT_WIDTH = 6    # 表が無い文書でも内容列をこれだけ確保

_ACCENT = "4C6EF5"
_LIGHT = "EEF2F7"
_CODE_BG = "F5F7FA"
_INK = "1F2933"
_BORDER = Side(style="thin", color="CBD2D9")
_TABLE_BORDER = Border(left=_BORDER, right=_BORDER, top=_BORDER, bottom=_BORDER)

_FONT_BODY = "Hiragino Kaku Gothic ProN"
_FONT_MONO = "Menlo"

_WRAP_TOP = Alignment(wrap_text=True, vertical="top")
_WRAP_MID = Alignment(wrap_text=True, vertical="center")


def _line_count(text: str, chars_per_line: int) -> int:
    total = 0
    for ln in text.split("\n"):
        total += max(1, -(-len(ln) // chars_per_line))  # ceil
    return total


def _key_tag(b: Block) -> str:
    if b.kind == "heading":
        return f"H{b.level}"
    if b.kind == "paragraph":
        return "P"
    if b.kind == "quote":
        return "Q"
    if b.kind == "code":
        return f"CODE:{b.lang}"
    if b.kind == "bullet":
        return f"UL:{b.level}"
    if b.kind == "ordered":
        return f"OL:{b.level}"
    if b.kind == "task":
        return f"CK:{1 if b.checked else 0}:{b.level}"
    raise ValueError(f"unexpected block kind for key tag: {b.kind}")


# ---------------------------------------------------------------------------
# ブロックモデル -> xlsx
# ---------------------------------------------------------------------------

def write_xlsx(blocks: list[Block], path: str, sheet_title: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Sheet1"

    content_width = max(
        DEFAULT_WIDTH,
        max((len(b.header) for b in blocks if b.kind == "table"), default=0),
    )
    last_col = CONTENT_COL + content_width - 1

    # 列幅と非表示キー列。
    ws.column_dimensions[get_column_letter(KEY_COL)].hidden = True
    ws.column_dimensions[get_column_letter(KEY_COL)].width = 3
    ws.column_dimensions[get_column_letter(CONTENT_COL)].width = 34
    for c in range(CONTENT_COL + 1, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18

    def merge_content(r: int) -> None:
        if last_col > CONTENT_COL:
            ws.merge_cells(
                start_row=r, start_column=CONTENT_COL, end_row=r, end_column=last_col
            )

    r = 1
    prev: Optional[Block] = None
    ord_counter = 0
    for b in blocks:
        if b.kind == "table":
            # 表ヘッダ
            ws.cell(r, KEY_COL, f"TBLH:{len(b.header)}")
            for j, h in enumerate(b.header):
                cell = ws.cell(r, CONTENT_COL + j, h)
                cell.font = Font(name=_FONT_BODY, bold=True, color=_INK)
                cell.fill = PatternFill("solid", fgColor=_LIGHT)
                cell.alignment = _WRAP_MID
                cell.border = _TABLE_BORDER
            r += 1
            # 表データ
            for row in b.rows:
                ws.cell(r, KEY_COL, f"TBLR:{len(b.header)}")
                for j in range(len(b.header)):
                    val = row[j] if j < len(row) else ""
                    cell = ws.cell(r, CONTENT_COL + j, val)
                    cell.font = Font(name=_FONT_BODY, color=_INK)
                    cell.alignment = _WRAP_TOP
                    cell.border = _TABLE_BORDER
                r += 1
            prev = b
            continue

        ws.cell(r, KEY_COL, _key_tag(b))
        cell = ws.cell(r, CONTENT_COL)
        merge_content(r)

        if b.kind == "heading":
            cell.value = b.text
            if b.level == 1:
                cell.font = Font(name=_FONT_BODY, bold=True, size=16, color=_INK)
                ws.row_dimensions[r].height = 26
            elif b.level == 2:
                cell.font = Font(name=_FONT_BODY, bold=True, size=12, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=_ACCENT)
                ws.row_dimensions[r].height = 22
            else:
                cell.font = Font(name=_FONT_BODY, bold=True, size=11, color=_INK)
                cell.fill = PatternFill("solid", fgColor=_LIGHT)
            cell.alignment = Alignment(vertical="center")
        elif b.kind == "paragraph":
            cell.value = b.text
            cell.font = Font(name=_FONT_BODY, color=_INK)
            cell.alignment = _WRAP_TOP
            ws.row_dimensions[r].height = 15 * _line_count(b.text, 56)
        elif b.kind == "quote":
            cell.value = b.text
            cell.font = Font(name=_FONT_BODY, italic=True, color="52606D")
            cell.fill = PatternFill("solid", fgColor=_LIGHT)
            cell.alignment = _WRAP_TOP
            ws.row_dimensions[r].height = 15 * _line_count(b.text, 56)
        elif b.kind == "code":
            cell.value = b.text
            cell.font = Font(name=_FONT_MONO, size=10, color=_INK)
            cell.fill = PatternFill("solid", fgColor=_CODE_BG)
            cell.alignment = _WRAP_TOP
            ws.row_dimensions[r].height = 15 * max(1, _line_count(b.text, 70))
        elif b.kind in ("bullet", "ordered", "task"):
            if b.kind == "task":
                prefix = "☑ " if b.checked else "☐ "
            elif b.kind == "ordered":
                # 連続する同レベルの番号付き項目は通し番号を振る（表示用）。
                if prev is not None and prev.kind == "ordered" and prev.level == b.level:
                    ord_counter += 1
                else:
                    ord_counter = 1
                prefix = f"{ord_counter}. "
            else:
                prefix = "• "
            cell.value = prefix + b.text
            cell.font = Font(name=_FONT_BODY, color=_INK)
            cell.alignment = Alignment(wrap_text=True, vertical="top", indent=b.level + 1)
            ws.row_dimensions[r].height = 15 * _line_count(b.text, 54)
        r += 1
        prev = b

    wb.save(path)


# ---------------------------------------------------------------------------
# xlsx -> ブロックモデル
# ---------------------------------------------------------------------------

def _cell_str(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def read_xlsx(path: str) -> tuple[list[Block], list[str]]:
    """xlsx を読み、(ブロック並び, 警告メッセージ) を返す。"""
    wb = load_workbook(path)
    ws = wb.active
    blocks: list[Block] = []
    warnings: list[str] = []
    max_col = ws.max_column

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        key = _cell_str(row[KEY_COL - 1].value).strip() if len(row) >= KEY_COL else ""
        content_cells = [_cell_str(c.value) for c in row[CONTENT_COL - 1:max_col]]
        first = content_cells[0] if content_cells else ""

        if key == "":
            # キーの無い行。内容があれば段落として拾い（行コピー以外で追記された行）、
            # 空行は読み飛ばす。
            if first.strip() != "":
                blocks.append(Block("paragraph", text=first))
                warnings.append(
                    f"行 {row[0].row}: 種別キーが無いため段落として取り込みました"
                    "（Excel で行を足すときは既存行をコピーしてください）。"
                )
            continue

        tag, _, arg = key.partition(":")

        if tag in ("H1", "H2", "H3", "H4", "H5", "H6"):
            blocks.append(Block("heading", text=first, level=int(tag[1:])))
        elif tag == "P":
            blocks.append(Block("paragraph", text=first))
        elif tag == "Q":
            blocks.append(Block("quote", text=first))
        elif tag == "CODE":
            blocks.append(Block("code", text=first, lang=arg))
        elif tag == "UL":
            blocks.append(Block("bullet", text=_strip_marker(first), level=int(arg or 0)))
        elif tag == "OL":
            blocks.append(Block("ordered", text=_strip_marker(first), level=int(arg or 0)))
        elif tag == "CK":
            checked_arg, _, level_arg = arg.partition(":")
            blocks.append(Block("task", text=_strip_marker(first),
                                checked=(checked_arg == "1"), level=int(level_arg or 0)))
        elif tag == "TBLH":
            ncols = int(arg or len(content_cells))
            blocks.append(Block("table", header=content_cells[:ncols], rows=[]))
        elif tag == "TBLR":
            ncols = int(arg or len(content_cells))
            if blocks and blocks[-1].kind == "table":
                blocks[-1].rows.append(content_cells[:ncols])
            else:
                warnings.append(f"行 {row[0].row}: 表ヘッダの無い表データ行を無視しました。")
        else:
            warnings.append(f"行 {row[0].row}: 未知の種別キー '{key}' を無視しました。")

    return blocks, warnings


_MARKER_RE = re.compile(r"^\s*(?:[•・*\-]|☑|☐|\d+\.)\s*")


def _strip_marker(text: str) -> str:
    """xlsx 表示用に付けた行頭マーカー（• / ☐ / ☑ 等）を取り除く。"""
    return _MARKER_RE.sub("", text, count=1)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
TEMPLATES = os.path.join(DOCS, "templates")

DOC_PAIRS = [
    ("要件定義書", os.path.join(DOCS, "requirements.md"), os.path.join(DOCS, "requirements.xlsx")),
    ("設計書", os.path.join(DOCS, "design.md"), os.path.join(DOCS, "design.xlsx")),
]

TEMPLATE_PAIRS = [
    ("要件定義書テンプレート",
     os.path.join(TEMPLATES, "requirements_template.md"),
     os.path.join(TEMPLATES, "requirements_template.xlsx")),
    ("設計書テンプレート",
     os.path.join(TEMPLATES, "design_template.md"),
     os.path.join(TEMPLATES, "design_template.xlsx")),
]


def _title_from(md_path: str, fallback: str) -> str:
    """Markdown の先頭 H1 をシート名に使う。無ければ fallback。"""
    try:
        with open(md_path, encoding="utf-8") as f:
            for line in f:
                m = _HEADING_RE.match(line.rstrip("\n"))
                if m and len(m.group(1)) == 1:
                    return m.group(2).strip()
    except OSError:
        pass
    return fallback


def cmd_export(md_path: str, xlsx_path: str, sheet_title: str = "") -> bool:
    if not os.path.exists(md_path):
        print(f"  skip: {os.path.relpath(md_path, ROOT)} が無いため変換しません。")
        return False
    with open(md_path, encoding="utf-8") as f:
        blocks = parse_markdown(f.read())
    title = sheet_title or _title_from(md_path, "ドキュメント")
    write_xlsx(blocks, xlsx_path, title)
    print(f"  export: {os.path.relpath(md_path, ROOT)} -> {os.path.relpath(xlsx_path, ROOT)}")
    return True


def cmd_import(xlsx_path: str, md_path: str) -> bool:
    if not os.path.exists(xlsx_path):
        print(f"  skip: {os.path.relpath(xlsx_path, ROOT)} が無いため取り込みません。")
        return False
    blocks, warnings = read_xlsx(xlsx_path)
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(render_markdown(blocks))
    print(f"  import: {os.path.relpath(xlsx_path, ROOT)} -> {os.path.relpath(md_path, ROOT)}")
    for w in warnings:
        print(f"    ! {w}")
    return True


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(
        description="要件定義書・設計書を Markdown ⇄ xlsx で双方向変換する。"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    p_exp = sub.add_parser("export", help="Markdown -> xlsx（単一ファイル）")
    p_exp.add_argument("input")
    p_exp.add_argument("output")

    p_imp = sub.add_parser("import", help="xlsx -> Markdown（単一ファイル）")
    p_imp.add_argument("input")
    p_imp.add_argument("output")

    sub.add_parser("export-docs", help="docs/{requirements,design}.md -> .xlsx")
    sub.add_parser("import-docs", help="docs/{requirements,design}.xlsx -> .md")
    sub.add_parser("template", help="docs/templates/*_template.md -> .xlsx")

    args = parser.parse_args(argv)

    if args.command == "export":
        return 0 if cmd_export(args.input, args.output) else 1
    if args.command == "import":
        return 0 if cmd_import(args.input, args.output) else 1
    if args.command == "export-docs":
        print("Markdown -> xlsx:")
        any_done = False
        for title, md, xlsx in DOC_PAIRS:
            any_done |= cmd_export(md, xlsx, title)
        if not any_done:
            print("  対象の docs/*.md がありません。先に TDD パイプラインで生成してください。")
        return 0
    if args.command == "import-docs":
        print("xlsx -> Markdown:")
        for _title, md, xlsx in DOC_PAIRS:
            cmd_import(xlsx, md)
        return 0
    if args.command == "template":
        print("テンプレート Markdown -> xlsx:")
        for title, md, xlsx in TEMPLATE_PAIRS:
            cmd_export(md, xlsx, title)
        return 0

    parser.print_help()
    return 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
